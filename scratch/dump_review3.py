import pptx
import json

prs = pptx.Presentation('reports/CyberMind AI – Project Review III.pptx')

slides_data = []
for i, slide in enumerate(prs.slides, 1):
    slide_info = {
        'slide_index': i,
        'shapes_count': len(slide.shapes),
        'texts': []
    }
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    font_name = p.font.name if p.font else None
                    font_size = p.font.size.pt if (p.font and p.font.size) else None
                    slide_info['texts'].append({
                        'text': txt,
                        'font_name': font_name,
                        'font_size': font_size,
                        'left': shape.left,
                        'top': shape.top,
                        'width': shape.width,
                        'height': shape.height
                    })
    slides_data.append(slide_info)

with open('scratch/review3_dump.json', 'w', encoding='utf-8') as f:
    json.dump(slides_data, f, indent=2, ensure_ascii=False)

print("Dumped review3 structure successfully!")
