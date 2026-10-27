import pptx

prs = pptx.Presentation(r"reports/CyberMind AI – Project Review IV.pptx")
print(f"Total slides in Review IV: {len(prs.slides)}")
print(f"Slide Dimensions: {prs.slide_width.inches:.2f}\" x {prs.slide_height.inches:.2f}\"")

for i, slide in enumerate(prs.slides, 1):
    title = "No Title"
    for shape in slide.shapes:
        if shape.has_text_frame and shape.top < 1000000:
            text = shape.text_frame.text.strip().replace('\n', ' ')
            if text:
                title = text
                break
    shapes_count = len(slide.shapes)
    print(f"Slide {i:2d} (shapes={shapes_count:2d}): {title}")
